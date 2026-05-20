"""
Импорт ежедневной выгрузки auto-parser.ru в SQLite.

Сканирует <inbox_dir> рекурсивно, по имени файла определяет:
  - platform: autoru / avito / drom → "Auto.ru" / "Avito Авто" / "Drom"
  - status: active / removed
  - vehicle_category: легковая / грузовик / автобус / прицеп / тягач / спецтехника / ...

Usage:
    python -m app.ingest <inbox_dir> [--limit-per-file N] [--snapshot YYYY-MM-DD]
"""
from __future__ import annotations

import argparse
import re
import sys
from datetime import date
from pathlib import Path

import openpyxl

from .db import get_conn, init_schema
from .normalize import normalize_brand, normalize_region

# Category recognition. The keys are substrings searched in the file's basename.
# More specific entries should come first.
CATEGORY_PATTERNS: list[tuple[str, str]] = [
    ("car_active", "Легковая"),
    ("car_removed", "Легковая"),
    ("lcv_", "Лёгкий коммерческий"),
    ("lyogkiy_kommercheskiy", "Лёгкий коммерческий"),
    ("sedelnye_tyagachi", "Тягач"),
    ("artic_", "Тягач"),
    ("polupritsepy", "Полуприцеп"),
    ("pritsepy", "Прицеп"),
    ("trailer_", "Прицеп"),
    ("gruzoviki", "Грузовик"),
    ("truck_", "Грузовик"),
    ("avtotsisterny", "Грузовик"),
    ("bus_", "Автобус"),
    ("avtobusy", "Автобус"),
    ("avtokrany", "Кран"),
    ("crane_", "Кран"),
    ("buldozery", "Бульдозер"),
    ("bulldozers", "Бульдозер"),
    ("ekskavatory", "Экскаватор"),
    ("dredge_", "Экскаватор"),
    ("kommunalnaya_tehnika", "Коммунальная техника"),
    ("municipal_", "Коммунальная техника"),
    ("dorozhno-stroitelnaya_tehnika", "Стройтехника"),
    ("construction_", "Стройтехника"),
    ("traktory_i_selhoztehnika", "Сельхозтехника"),
    ("agricultural_", "Сельхозтехника"),
    ("mini-traktory", "Сельхозтехника"),
    ("traktory_kommunalnye", "Сельхозтехника"),
    ("traktory", "Сельхозтехника"),
    ("pogruzchiki", "Погрузчик"),
    ("autoloader_", "Погрузчик"),
    ("vilochnye_pogruzchiki", "Погрузчик"),
    ("teleskopicheskie_pogruzchiki", "Погрузчик"),
    ("shtabelyory", "Погрузчик"),
    ("avtodoma", "Дом на колёсах"),
    ("vezdehody", "Вездеход"),
    ("lesozagotovitelnaya_tehnika", "Лесозаготовительная"),
    ("tehnika_dlya_lesozagotovki", "Лесозаготовительная"),
    ("drugaya_spetstehnika", "Прочая спецтехника"),
    ("polivomoechnye_mashiny", "Поливо-моечная"),
    ("podmetalno-uborochnye_mashiny", "Подметально-уборочная"),
    ("kombinirovannye_dorozhnye_mashiny", "Комбинированная дорожная"),
    ("seyalki", "Сеялка"),
    ("press-podborschiki", "Пресс-подборщик"),
    ("plugi", "Плуг"),
    ("borony", "Борона"),
    ("kultivatory", "Культиватор"),
    ("kosilki", "Косилка"),
    ("opryskivateli", "Опрыскиватель"),
    ("pochvofrezy", "Почвофреза"),
    ("grabli_voroshilki", "Грабли"),
    ("navesnoe_oborudovanie", "Навесное оборудование"),
    ("bunkerovozy", "Бункеровоз"),
    ("assenizatory", "Ассенизатор"),
]

PLATFORM_BY_PREFIX = {
    "autoru": "Auto.ru",
    "avito": "Avito Авто",
    "drom": "Drom",
}

COLUMNS = [
    "inner_id", "platform", "status", "snapshot_date", "vehicle_category",
    "mark", "model", "generation", "configuration", "complectation",
    "body_type", "year", "price_rub", "km_age", "engine_type",
    "displacement", "horse_power", "transmission", "drive_type",
    "color", "wheel", "section", "condition", "owners_count",
    "no_accidents", "seller", "seller_type", "seller_id", "seller_url",
    "region", "city", "address", "url",
    "vin", "pts", "custom",
    "in_stock", "state_not_beaten",
    "price_usd", "price_eur",
    "offer_updated_at", "image_urls",
    "options", "description",
]


def _to_int(v):
    try:
        if v is None or v == "":
            return None
        return int(float(v))
    except (ValueError, TypeError):
        return None


def _to_float(v):
    try:
        if v is None or v == "":
            return None
        return float(v)
    except (ValueError, TypeError):
        return None


def _to_str(v):
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    return s if s and s.lower() != "none" else None


def _detect(filename: str) -> tuple[str | None, str | None, str | None]:
    """Returns (platform, status, category) or Nones if not recognised."""
    name = filename.lower()
    platform: str | None = None
    for prefix, label in PLATFORM_BY_PREFIX.items():
        if name.startswith(prefix + "_"):
            platform = label
            break
    if not platform:
        return None, None, None
    status = "removed" if "_removed" in name or "removed." in name else "active"
    category: str | None = None
    for needle, label in CATEGORY_PATTERNS:
        if needle in name:
            category = label
            break
    return platform, status, category


def _record(row: dict, platform: str, status: str, snapshot: str, category: str | None) -> dict:
    return {
        "inner_id": _to_str(row.get("inner_id") or row.get("id")),
        "platform": platform,
        "status": status,
        "snapshot_date": snapshot,
        "vehicle_category": category,
        "mark": normalize_brand(_to_str(row.get("mark"))),
        "model": _to_str(row.get("model")),
        "generation": _to_str(row.get("generation")),
        "configuration": _to_str(row.get("configuration")),
        "complectation": _to_str(row.get("complectation")),
        "body_type": _to_str(row.get("body_type")),
        "year": _to_int(row.get("year")),
        "price_rub": _to_int(row.get("price_rub")),
        "km_age": _to_int(row.get("km_age")),
        "engine_type": _to_str(row.get("engine_type")),
        "displacement": _to_float(row.get("displacement")),
        "horse_power": _to_int(row.get("horse_power")),
        "transmission": _to_str(row.get("transmission")),
        "drive_type": _to_str(row.get("drive_type")),
        "color": _to_str(row.get("color")),
        "wheel": _to_str(row.get("wheel")),
        "section": _to_str(row.get("section")),
        "condition": _to_str(row.get("condition")),
        "owners_count": _to_int(row.get("owners_count")),
        "no_accidents": _to_str(row.get("no_accidents")),
        "seller": _to_str(row.get("seller")),
        "seller_type": _to_str(row.get("seller_type")),
        "seller_id": _to_str(row.get("seller_id")),
        "seller_url": _to_str(row.get("seller_url")),
        "region": normalize_region(_to_str(row.get("region"))),
        "city": _to_str(row.get("city")),
        "address": _to_str(row.get("address")),
        "url": _to_str(row.get("url")),
        "vin": _to_str(row.get("vin")),
        "pts": _to_str(row.get("pts")),
        "custom": _to_str(row.get("custom")),
        "in_stock": _to_str(row.get("in_stock")),
        "state_not_beaten": _to_str(row.get("state_not_beaten")),
        "price_usd": _to_int(row.get("price_usd")),
        "price_eur": _to_int(row.get("price_eur")),
        "offer_updated_at": _to_str(row.get("offer_updated_at")),
        "image_urls": _to_str(row.get("image_urls")),
        "options": _to_str(row.get("options")),
        "description": _to_str(row.get("description")),
    }


def _import_file(
    conn,
    path: Path,
    platform: str,
    status: str,
    category: str | None,
    snapshot: str,
    limit: int | None,
) -> int:
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as e:
        print(f"  [skip] {path.name}: {e}")
        return 0
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    try:
        header = list(next(rows))
    except StopIteration:
        wb.close()
        return 0

    batch = []
    BATCH = 500
    inserted = 0
    placeholders = ",".join("?" * len(COLUMNS))
    sql = f"INSERT OR IGNORE INTO listings ({','.join(COLUMNS)}) VALUES ({placeholders})"

    for row in rows:
        if limit is not None and inserted >= limit:
            break
        rec = dict(zip(header, row))
        r = _record(rec, platform, status, snapshot, category)
        if not r["inner_id"] or not r["mark"]:
            continue
        batch.append(tuple(r[c] for c in COLUMNS))
        inserted += 1
        if len(batch) >= BATCH:
            conn.executemany(sql, batch)
            batch.clear()
    if batch:
        conn.executemany(sql, batch)
    wb.close()
    return inserted


def ingest(inbox: Path, limit_per_file: int | None = None, snapshot: str | None = None) -> None:
    snapshot = snapshot or date.today().isoformat()
    init_schema()
    conn = get_conn()
    conn.execute("DELETE FROM listings")

    files = sorted(p for p in inbox.rglob("*.xlsx") if not p.name.startswith("~$"))
    print(f"Found {len(files)} xlsx files")

    totals: dict[str, int] = {}
    try:
        for path in files:
            platform, status, category = _detect(path.name)
            if not platform:
                print(f"  [unknown] {path.name}")
                continue
            n = _import_file(conn, path, platform, status, category, snapshot, limit_per_file)
            key = f"{platform} / {category or '—'} / {status}"
            totals[key] = totals.get(key, 0) + n
            print(f"  {path.name}: {n}")

        conn.execute("INSERT OR REPLACE INTO meta(key, value) VALUES('last_snapshot', ?)", (snapshot,))
        conn.commit()
        print(f"\nBy bucket:")
        for k, v in sorted(totals.items()):
            print(f"  {k:60} {v}")
        print(f"\nTotal: {sum(totals.values())} rows (snapshot={snapshot})")
    finally:
        conn.close()


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("inbox", help="Directory with xlsx files (scanned recursively)")
    parser.add_argument("--limit-per-file", type=int, default=None)
    parser.add_argument("--snapshot", default=None, help="YYYY-MM-DD; defaults to today")
    args = parser.parse_args()
    inbox = Path(args.inbox)
    if not inbox.exists():
        print(f"Inbox not found: {inbox}", file=sys.stderr)
        return 1
    ingest(inbox, args.limit_per_file, args.snapshot)
    return 0


if __name__ == "__main__":
    sys.exit(main())
